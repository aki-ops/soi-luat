from __future__ import annotations

from pathlib import Path

from docx import Document

from soiluat.cards import OCR_LABEL, Card

HIEU_LUC = {"con": "Còn hiệu lực", "het": "Hết hiệu lực", "chua_ro": "Hiệu lực chưa rõ"}  # như _results.html


def page_label(page_from, page_to=None) -> str:
    """"Trang x" hoặc "Trang x–y"; rỗng khi chưa biết trang."""
    if not page_from:
        return ""
    if page_to and page_to != page_from:
        return f"Trang {page_from}–{page_to}"
    return f"Trang {page_from}"


def card_pages(c: Card) -> str:
    # page_to / notes / intent do agent C thêm vào Card — đọc phòng thủ
    return page_label(getattr(c, "page_from", None), getattr(c, "page_to", None))


def card_notes(c: Card) -> list[str]:
    return [n for n in (getattr(c, "notes", None) or []) if n][:3]


def export_cards(cards: list[Card], dest: Path) -> None:
    doc = Document()
    doc.add_heading("Báo cáo Soi Luật", level=0)
    for c in cards:
        doc.add_heading(c.kind + " — " + (c.title or ""), level=1)
        addr = c.duty_addr or c.sanction_addr or c.so_hieu
        meta = [f"Hạng {c.rank}", addr, card_pages(c), c.file_path, HIEU_LUC.get(c.hieu_luc, c.hieu_luc)]
        doc.add_paragraph(" · ".join(x for x in meta if x))
        if c.sample:
            doc.add_paragraph(c.sample)
        if c.duty_text:
            p = doc.add_paragraph(c.duty_text)
            p.add_run(f"\nNguồn: {c.duty_addr}")
        if c.sanction_text:
            p = doc.add_paragraph(c.sanction_text)
            p.add_run(f"\nNguồn: {c.sanction_addr}")
        for n in card_notes(c):
            doc.add_paragraph(f"Ghi chú: {n}")
        warns = list(c.warnings)
        if c.ocr_unverified and OCR_LABEL not in warns:
            warns.append(OCR_LABEL)
        for w in warns:
            doc.add_paragraph(f"Cảnh báo: {w}")
    dest.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(dest))

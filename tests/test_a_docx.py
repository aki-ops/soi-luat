"""L1: DOCX theo thứ tự thân văn bản — bảng, footnote/endnote, comment, text box, header/footer, ngắt trang."""

import shutil
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import parse_xml

from soiluat.extract import extract_docx

W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
MC = 'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"'


def _notes_xml(tag: str, items: dict[int, str]) -> str:
    body = "".join(
        f'<w:{tag} w:id="{i}"><w:p><w:r><w:{tag}Ref/></w:r><w:r><w:t xml:space="preserve"> {t}</w:t></w:r></w:p></w:{tag}>'
        for i, t in items.items()
    )
    sep = f'<w:{tag} w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:{tag}>'
    return f'<?xml version="1.0" encoding="UTF-8"?><w:{tag}s {W}>{sep}{body}</w:{tag}s>'


def _add_parts(path: Path, parts: dict[str, str]) -> None:
    """Chép thêm part vào gói docx (python-docx không tự tạo footnotes/comments)."""
    tmp = path.with_suffix(".tmp")
    shutil.move(path, tmp)
    with zipfile.ZipFile(tmp) as zin, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            zout.writestr(item, zin.read(item.filename))
        for name, xml in parts.items():
            zout.writestr(name, xml.encode("utf-8"))
    tmp.unlink()


def _textbox_run(text: str) -> str:
    box = f"<w:txbxContent><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:txbxContent>"
    return (
        f"<w:r {W} {MC}><mc:AlternateContent>"
        f"<mc:Choice Requires=\"wps\"><w:drawing><wps:txbx xmlns:wps=\"http://schemas.microsoft.com/office/word/2010/wordprocessingShape\">{box}</wps:txbx></w:drawing></mc:Choice>"
        f"<mc:Fallback><w:pict><v:textbox xmlns:v=\"urn:schemas-microsoft-com:vml\">{box}</v:textbox></w:pict></mc:Fallback>"
        "</mc:AlternateContent></w:r>"
    )


def _build(tmp_path: Path) -> Path:
    d = Document()
    d.sections[0].header.paragraphs[0].text = "BẢN LƯU NỘI BỘ"
    d.sections[0].footer.paragraphs[0].text = "Đã rà soát"
    d.add_paragraph("Điều 1. Phạm vi")
    p = d.add_paragraph("Văn bản này sửa đổi")
    p._p.append(parse_xml(f'<w:r {W}><w:footnoteReference w:id="1"/></w:r>'))
    p._p.append(parse_xml(f'<w:r {W}><w:t xml:space="preserve"> theo luật mới.</w:t></w:r>'))
    d.add_paragraph("Mức phạt như bảng sau:")
    t = d.add_table(rows=3, cols=3)
    for j, v in enumerate(["STT", "Hành vi", "Mức phạt"]):
        t.cell(0, j).text = v
    t.cell(1, 0).text = "1"
    t.cell(1, 1).text = "Không niêm yết giá"
    t.cell(1, 2).text = "Phạt tiền 3 triệu"
    t.cell(2, 0).text = "2"
    t.cell(1, 2).merge(t.cell(2, 2))  # vMerge
    t.cell(2, 1).text = "Không lập sổ"
    d.add_paragraph("Sau bảng là đoạn này.")
    c = d.add_paragraph()
    c._p.append(parse_xml(f'<w:commentRangeStart {W} w:id="0"/>'))
    c._p.append(parse_xml(f'<w:r {W}><w:t>Điều 2. Đối tượng áp dụng</w:t></w:r>'))
    c._p.append(parse_xml(f'<w:commentRangeEnd {W} w:id="0"/>'))
    c._p.append(parse_xml(_textbox_run("Ghi chú trong khung")))
    br = d.add_paragraph("Hết trang một")
    br.runs[0].add_break(WD_BREAK.PAGE)
    # Word thường ghi thêm lastRenderedPageBreak ngay sau ngắt cứng: không được đẻ trang rỗng
    p3 = d.add_paragraph()
    p3._p.append(parse_xml(f'<w:r {W}><w:lastRenderedPageBreak/><w:t>Điều 3. Trang hai</w:t></w:r>'))
    e = d.add_paragraph("Câu có chú thích cuối")
    e._p.append(parse_xml(f'<w:r {W}><w:endnoteReference w:id="2"/></w:r>'))
    p4 = d.add_paragraph()
    p4._p.append(parse_xml(f'<w:r {W}><w:t>Nửa đầu</w:t><w:lastRenderedPageBreak/><w:t>Điều 4. Trang ba</w:t></w:r>'))
    path = tmp_path / "f.docx"
    d.save(str(path))
    comments = (
        f'<?xml version="1.0" encoding="UTF-8"?><w:comments {W}><w:comment w:id="0" w:author="A">'
        "<w:p><w:r><w:t>Cần đối chiếu Thông tư 45</w:t></w:r></w:p></w:comment></w:comments>"
    )
    _add_parts(
        path,
        {
            "word/footnotes.xml": _notes_xml("footnote", {1: "Sửa bởi Nghị định 123/2024/NĐ-CP."}),
            "word/endnotes.xml": _notes_xml("endnote", {2: "Chú thích cuối văn bản."}),
            "word/comments.xml": comments,
        },
    )
    return path


def test_docx_body_order_tables_notes_pages(tmp_path):
    pages = extract_docx(_build(tmp_path))
    assert [p.page_no for p in pages] == [1, 2, 3]
    t1 = pages[0].text.splitlines()
    # bảng nội tuyến đúng vị trí, ô gộp dọc giữ chỗ rỗng
    i = t1.index("Mức phạt như bảng sau:")
    assert t1[i + 1 : i + 5] == [
        "| STT | Hành vi | Mức phạt |",
        "| 1 | Không niêm yết giá | Phạt tiền 3 triệu |",
        "| 2 | Không lập sổ |  |",
        "Sau bảng là đoạn này.",
    ]
    assert "Văn bản này sửa đổi[^1] theo luật mới." in t1
    kinds = {n.kind: n for n in pages[0].notes}
    assert kinds["footnote"].ref == "1" and "123/2024" in kinds["footnote"].text
    assert kinds["comment"].ref == "Điều 2. Đối tượng áp dụng"
    assert "Thông tư 45" in kinds["comment"].text
    assert kinds["header"].text == "BẢN LƯU NỘI BỘ"
    assert kinds["footer"].text == "Đã rà soát"
    # text box có cả Choice và Fallback: chỉ một ghi chú, không vào thân
    assert [n.text for n in pages[0].notes if n.kind == "textbox"] == ["Ghi chú trong khung"]
    assert "Ghi chú trong khung" not in pages[0].text
    assert pages[1].text.startswith("Điều 3. Trang hai")
    assert "Câu có chú thích cuối[^2]" in pages[1].text
    assert [(n.kind, n.ref) for n in pages[1].notes] == [("endnote", "2")]
    assert pages[1].text.endswith("Nửa đầu")
    assert pages[2].text == "Điều 4. Trang ba"
    assert pages[2].notes == []


def test_docx_grid_span_keeps_columns(tmp_path):
    d = Document()
    t = d.add_table(rows=2, cols=3)
    t.cell(0, 0).merge(t.cell(0, 1)).text = "Gộp ngang"
    t.cell(0, 2).text = "C"
    for j, v in enumerate("xyz"):
        t.cell(1, j).text = v
    path = tmp_path / "g.docx"
    d.save(str(path))
    assert extract_docx(path)[0].text.splitlines() == ["| Gộp ngang |  | C |", "| x | y | z |"]


def test_docx_numbering_in_body_walk(tmp_path):
    d = Document()
    d.add_paragraph("Điều 4. Báo cáo")
    d.add_paragraph("Gửi báo cáo định kỳ.", style="List Number")
    d.add_paragraph("Lưu hồ sơ.", style="List Number")
    path = tmp_path / "n.docx"
    d.save(str(path))
    assert extract_docx(path)[0].text.splitlines() == [
        "Điều 4. Báo cáo",
        "1. Gửi báo cáo định kỳ.",
        "2. Lưu hồ sơ.",
    ]

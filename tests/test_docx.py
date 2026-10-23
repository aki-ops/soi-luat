from pathlib import Path

from docx import Document

from soiluat.extract import extract_docx
from soiluat.ingest import ingest_path


def test_ingest_docx(db, tmp_path):
    d = Document()
    d.add_paragraph("Nghị định số 99/2021/NĐ-CP")
    d.add_paragraph("Điều 2. Nghĩa vụ báo cáo")
    d.add_paragraph("1. Doanh nghiệp phải gửi báo cáo định kỳ.")
    path = tmp_path / "a.docx"
    d.save(str(path))
    doc_id = ingest_path(db, path, tmp_path)
    n = db.execute(
        "SELECT COUNT(*) AS n FROM units WHERE doc_id = ?", (doc_id,)
    ).fetchone()["n"]
    assert n >= 2
    labs = db.execute(
        "SELECT COUNT(*) AS n FROM unit_labels WHERE label = 'NGHIA_VU'"
    ).fetchone()["n"]
    assert labs >= 1


def test_docx_list_number_prefix(tmp_path):
    d = Document()
    d.add_paragraph("Điều 4. Báo cáo")
    d.add_paragraph("Doanh nghiệp phải gửi báo cáo định kỳ.", style="List Number")
    path = tmp_path / "n.docx"
    d.save(str(path))
    text = extract_docx(path)[0].text
    assert "Điều 4" in text
    assert "1." in text or "1)" in text
